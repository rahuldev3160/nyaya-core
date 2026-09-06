from docx import Document


def extract_text(filepath: str) -> list[tuple[int, str]]:
    """Word has no fixed page concept before rendering — returns a single (1, text)
    page so the chunker's interface stays uniform across all 7 parsers (BUG-02)."""
    doc = Document(filepath)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
            if row_text:
                parts.append(row_text)
    return [(1, "\n\n".join(parts))]
